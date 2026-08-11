"""Renderer tests for the RAIS register adapter.

The API tests cover the happy path end to end; these pin the two ways the object
table can depart from the blank form's fifteen printed lines, plus the empty
conclusion date.
"""

import io
from datetime import UTC, date, datetime

from docx import Document as DocxDocument

from app.use_of_collections.application.documents import (
    ObjectAccessLogDocument,
    ObjectAccessLogDocumentObject,
)
from app.use_of_collections.infrastructure.object_access_log_docx import (
    DocxObjectAccessLogRenderer,
)

_HEADER_ROWS = 1
_FORM_OBJECT_LINES = 15


def _document(
    objects: tuple[ObjectAccessLogDocumentObject, ...],
    conclusion_date: date | None = None,
) -> ObjectAccessLogDocument:
    return ObjectAccessLogDocument(
        reference_number="OAL-ABCDEFG1",
        issued_on=date(2026, 8, 11),
        requester="Ana Silva",
        researcher_name="Ana Silva",
        researcher_email="ana@example.org",
        collection="Zoologia",
        curator="Nuno Curador",
        conclusion_date=conclusion_date,
        objects=objects,
    )


def _object(inventory_number: str) -> ObjectAccessLogDocumentObject:
    return ObjectAccessLogDocumentObject(
        inventory_number=inventory_number,
        designation="Vulpes vulpes",
        object_type="peles",
        number_of_objects=1,
        accessed_at=datetime(2026, 6, 2, 14, 30, tzinfo=UTC),
        observations="",
    )


async def _render(document: ObjectAccessLogDocument) -> DocxDocument:
    content = await DocxObjectAccessLogRenderer().render(document)
    return DocxDocument(io.BytesIO(content))


async def test_render_grows_the_object_table_past_the_printed_lines() -> None:
    objects = tuple(_object(f"INV-{i:03d}") for i in range(1, 21))

    rendered = await _render(_document(objects))

    table = rendered.tables[2]
    assert len(table.rows) == _HEADER_ROWS + 20
    assert [row.cells[0].text for row in table.rows[1:]] == [
        obj.inventory_number for obj in objects
    ]


async def test_render_keeps_the_printed_lines_when_the_log_is_empty() -> None:
    rendered = await _render(_document(()))

    table = rendered.tables[2]
    assert len(table.rows) == _HEADER_ROWS + _FORM_OBJECT_LINES
    assert all(cell.text == "" for row in table.rows[1:] for cell in row.cells)


async def test_render_leaves_the_conclusion_date_blank_while_open() -> None:
    rendered = await _render(_document((_object("INV-001"),)))

    assert rendered.tables[3].rows[1].cells[1].text == ""


async def test_render_prints_dates_in_the_forms_day_first_format() -> None:
    rendered = await _render(
        _document((_object("INV-001"),), conclusion_date=date(2026, 6, 9))
    )

    assert rendered.tables[0].rows[2].cells[2].text == "11-08-2026"
    assert rendered.tables[2].rows[1].cells[4].text == "02-06-2026"
    assert rendered.tables[3].rows[1].cells[1].text == "09-06-2026"
