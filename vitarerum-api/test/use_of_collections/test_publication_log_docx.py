"""Renderer tests for the RRP publication/output register."""

import io
import zipfile
from datetime import UTC, date, datetime

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from app.use_of_collections.application.documents import (
    PublicationLogDocument,
    PublicationLogDocumentAttachment,
    PublicationLogDocumentEntry,
)
from app.use_of_collections.infrastructure.publication_log_docx import (
    DocxPublicationLogRenderer,
)


def _entry(
    sequence: int,
    note: str,
    *,
    attachments: tuple[PublicationLogDocumentAttachment, ...] = (),
) -> PublicationLogDocumentEntry:
    return PublicationLogDocumentEntry(
        sequence=sequence,
        added_at=datetime(2026, 6, sequence, 14, 30, tzinfo=UTC),
        added_by="Ana Silva",
        note=note,
        object_reference="INV-001 — Tursiops truncatus\nZoologia",
        attachments=attachments,
    )


def _document(*entries: PublicationLogDocumentEntry) -> PublicationLogDocument:
    return PublicationLogDocument(
        reference_number="PUB-MUHNAC/COL/2026/0001",
        issued_on=date(2026, 8, 30),
        project_reference="PRJ-MUHNAC/COL/2026/0001",
        project_title="Estudo da coleção zoológica",
        requester="Ana Silva",
        curator="Nuno Curador",
        entries=entries,
    )


async def _render(*entries: PublicationLogDocumentEntry) -> bytes:
    return await DocxPublicationLogRenderer().render(_document(*entries))


async def test_render_repeats_one_chronological_row_per_publication() -> None:
    content = await _render(_entry(1, "Primeiro artigo"), _entry(2, "Segundo artigo"))

    rendered = DocxDocument(io.BytesIO(content))
    rows = rendered.tables[2].rows
    assert len(rows) == 3
    assert [row.cells[0].text for row in rows[1:]] == ["1", "2"]
    assert [row.cells[3].text for row in rows[1:]] == [
        "Primeiro artigo",
        "Segundo artigo",
    ]
    assert rows[1].cells[1].text == "01-06-2026"


async def test_render_names_attachments_one_per_line_and_marks_an_empty_list() -> None:
    content = await _render(
        _entry(
            1,
            "Com ficheiros",
            attachments=(
                PublicationLogDocumentAttachment("paper.pdf", "Artigo aceite"),
                PublicationLogDocumentAttachment("poster.png", ""),
            ),
        ),
        _entry(2, "Sem ficheiros"),
    )

    rendered = DocxDocument(io.BytesIO(content))
    rows = rendered.tables[2].rows
    assert rows[1].cells[5].text == "paper.pdf — Artigo aceite\nposter.png"
    assert rows[2].cells[5].text == "—"

    body = zipfile.ZipFile(io.BytesIO(content)).read("word/document.xml").decode()
    start = body.index("paper.pdf")
    end = body.index("poster.png")
    assert "<w:br/>" in body[start:end]


async def test_render_repeats_the_table_header_on_later_pages() -> None:
    content = await _render(
        *(_entry(index, f"Publicação {index}") for index in range(1, 30))
    )

    header = DocxDocument(io.BytesIO(content)).tables[2].rows[0]
    properties = header._tr.find(qn("w:trPr"))
    assert properties is not None
    assert properties.find(qn("w:tblHeader")) is not None
    assert properties.find(qn("w:cantSplit")) is not None

    for row in DocxDocument(io.BytesIO(content)).tables[2].rows[1:]:
        row_properties = row._tr.find(qn("w:trPr"))
        assert row_properties is not None
        assert row_properties.find(qn("w:cantSplit")) is not None
