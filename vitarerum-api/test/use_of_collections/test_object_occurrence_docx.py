"""Renderer tests for the ROC report adapter.

The API tests cover the happy path end to end; these pin the parts of the form
that carry free text — multi-line descriptions have to come out as real line
breaks, not as a single run with a stray newline Word would swallow.
"""

import io
import zipfile
from datetime import UTC, date, datetime

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from app.use_of_collections.application.documents import (
    ObjectOccurrenceDocument,
    ObjectOccurrenceDocumentImage,
)
from app.use_of_collections.infrastructure.object_occurrence_docx import (
    DocxObjectOccurrenceRenderer,
)


def _document(
    *,
    reference_number: str = "OOL-ABCDEFG1",
    number_of_objects: int = 1,
    detailed_description: str = "Dano observado na base.",
    testimonial: str = "",
    images: tuple[ObjectOccurrenceDocumentImage, ...] = (),
) -> ObjectOccurrenceDocument:
    return ObjectOccurrenceDocument(
        reference_number=reference_number,
        issued_on=date(2026, 8, 11),
        institution="Ana Silva",
        collection="Zoologia",
        designation="Tursiops truncatus",
        inventory_number="INV-001",
        number_of_objects=number_of_objects,
        occurred_at=datetime(2026, 6, 2, 14, 30, tzinfo=UTC),
        location="Sala 3, reserva",
        detailed_description=detailed_description,
        testimonial=testimonial,
        images=images,
        reported_by="Nuno Curador",
    )


async def _render(document: ObjectOccurrenceDocument) -> bytes:
    return await DocxObjectOccurrenceRenderer().render(document)


def _fields(content: bytes) -> dict[str, str]:
    rendered = DocxDocument(io.BytesIO(content))
    return {
        row.cells[0].text.split("\n")[0].strip(): row.cells[1].text
        for row in rendered.tables[1].rows
    }


async def test_render_breaks_a_multi_line_description_into_line_breaks() -> None:
    content = await _render(_document(detailed_description="Linha A.\nLinha B."))

    body = zipfile.ZipFile(io.BytesIO(content)).read("word/document.xml").decode()
    start = body.index("Linha A.")
    end = body.index("Linha B.")
    assert "<w:br/>" in body[start:end]


async def test_render_lets_the_header_row_grow_for_a_full_reference() -> None:
    """A reference under the active mask wraps at its hyphen; the row it sits in
    must be free to grow, or Word clips everything after the first line and only
    "OO-" reaches the page."""
    content = await _render(_document(reference_number="OO-MUHNAC/COL/2026/0001"))

    rendered = DocxDocument(io.BytesIO(content))
    header_row = rendered.tables[0].rows[2]
    assert header_row.cells[0].text == "OO-MUHNAC/COL/2026/0001"
    height = header_row._tr.find(qn("w:trPr")).find(qn("w:trHeight"))
    assert height.get(qn("w:hRule")) != "exact"


async def test_render_omits_the_quantity_for_a_single_object() -> None:
    content = await _render(_document(number_of_objects=1))

    assert _fields(content)["Nº(s) Inventário"] == "INV-001"


async def test_render_names_the_attached_images_one_per_line() -> None:
    content = await _render(
        _document(
            images=(
                ObjectOccurrenceDocumentImage("before.jpg", "Antes"),
                ObjectOccurrenceDocumentImage("after.jpg", ""),
            )
        )
    )

    assert _fields(content)["Imagens"] == "before.jpg — Antes\nafter.jpg"


async def test_render_leaves_optional_fields_empty_rather_than_instructional() -> None:
    content = await _render(_document(testimonial="", images=()))

    fields = _fields(content)
    assert fields["Depoimentos recolhidos (se aplicável)"] == ""
    assert fields["Imagens"] == ""
