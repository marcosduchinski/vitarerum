"""DB-backed coverage for the collection data-source management flow.

Exercises the SQLAlchemy repositories/index and the use cases over a real
(in-memory SQLite) session, with a real .xlsx parsed by openpyxl and an
in-memory file storage.
"""

import io
from datetime import UTC, datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from app.collection_object_index.application.use_cases import (
    AssignCollectionCurator,
    CreateCollection,
    CreateCollectionArea,
    CreateCollectionAreaInput,
    CreateCollectionInput,
    DeleteSourceDocument,
    GetCollection,
    ListCollectionAreas,
    ListCollectionSourceDocuments,
    ListCuratorCandidates,
    ListManageableCollections,
    ListSearchableCollections,
    MoveCollectionToArea,
    MoveCollectionToAreaInput,
    ReindexSourceDocument,
    RemoveCollection,
    RemoveCollectionArea,
    RemoveCollectionCurator,
    UpdateCollection,
    UpdateCollectionArea,
    UpdateCollectionAreaInput,
    UpdateCollectionInput,
    UploadSourceDocument,
    UploadSourceDocumentInput,
)
from app.collection_object_index.domain.enums import SourceDocumentStatus
from app.collection_object_index.domain.models import (
    CollectionAreaId,
    CollectionAreaInUse,
    CollectionAreaNotFound,
    CollectionId,
    CollectionNotFound,
)
from app.collection_object_index.infrastructure.models import (
    CollectionAreaRecord,
    CollectionObjectRecord,
    CollectionRecord,
)
from app.collection_object_index.infrastructure.parser_openpyxl import (
    OpenpyxlCollectionObjectParser,
)
from app.collection_object_index.infrastructure.repositories import (
    SqlAlchemyCollectionObjectIndex,
    SqlAlchemyCollectionRepository,
    SqlAlchemySourceDocumentRepository,
)
from app.database import Base
from app.identity.application.read_models import PermissionView, UserView
from app.identity.public import Actor, GroupName, PermissionId
from app.shared.exceptions import AccessDenied, InsufficientGroup

_NOW = datetime(2026, 7, 4, 12, 0, tzinfo=UTC)
_REPTILES = CollectionId("col-zoo")
_FISH = CollectionId("col-bot")
_ZOOLOGY_AREA = CollectionAreaId("area-nat-hist")
_TEMPORARY_AREA = CollectionAreaId("area-hum-sci")

_ADMIN = Actor(
    id=PermissionId("perm-admin"), group=GroupName.SYS_ADMIN, email="a@museum.pt"
)
_MANAGER = Actor(
    id=PermissionId("perm-mgmt"),
    group=GroupName.COLLECTIONS_MANAGEMENT,
    email="m@museum.pt",
)
_CURATOR = Actor(
    id=PermissionId("perm-cur"), group=GroupName.CURATORIAL, email="c@museum.pt"
)


class _FakeClock:
    def now(self) -> datetime:
        return _NOW


class _FakeStorage:
    def __init__(self) -> None:
        self.saved: dict[str, bytes] = {}
        self.deleted: list[str] = []

    async def save(self, content: bytes, file_reference: str) -> str:
        self.saved[file_reference] = content
        return file_reference

    async def read(self, file_reference: str) -> bytes:
        return self.saved[file_reference]

    async def delete(self, file_reference: str) -> None:
        self.deleted.append(file_reference)
        self.saved.pop(file_reference, None)


def _xlsx(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Objects"
    sheet.append(["Inventory No", "Name"])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _session_factory() -> async_sessionmaker:
    engine = create_async_engine(
        "sqlite+aiosqlite://",
        poolclass=StaticPool,
        connect_args={"check_same_thread": False},
    )
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(bind=engine, expire_on_commit=False)
    async with factory() as session:
        session.add_all(
            [
                CollectionAreaRecord(
                    id=str(_ZOOLOGY_AREA),
                    name="Zoology",
                    created_at=_NOW,
                    updated_at=_NOW,
                ),
                CollectionAreaRecord(
                    id=str(_TEMPORARY_AREA),
                    name="Temporary Area",
                    created_at=_NOW,
                    updated_at=_NOW,
                ),
            ]
        )
        session.add_all(
            [
                CollectionRecord(
                    id=str(_REPTILES),
                    area_id=str(_ZOOLOGY_AREA),
                    name="REPTILES & AMPHIBIANS",
                    created_at=_NOW,
                    updated_at=_NOW,
                ),
                CollectionRecord(
                    id=str(_FISH),
                    area_id=str(_ZOOLOGY_AREA),
                    name="FISH",
                    created_at=_NOW,
                    updated_at=_NOW,
                ),
            ]
        )
        await session.commit()
    return factory


def _uploader(session: AsyncSession, storage: _FakeStorage) -> UploadSourceDocument:
    return UploadSourceDocument(
        SqlAlchemyCollectionRepository(session),
        SqlAlchemySourceDocumentRepository(session),
        storage,
        OpenpyxlCollectionObjectParser(),
        SqlAlchemyCollectionObjectIndex(session),
        _FakeClock(),
    )


async def _object_rows(session: AsyncSession) -> list[CollectionObjectRecord]:
    result = await session.execute(
        select(CollectionObjectRecord).order_by(CollectionObjectRecord.row_number)
    )
    return list(result.scalars())


async def _upload(
    factory: async_sessionmaker,
    storage: _FakeStorage,
    caller: Actor = _ADMIN,
    collection_id: CollectionId = _REPTILES,
    file_name: str = "zoo.xlsx",
    rows: list[list[str]] | None = None,
):
    async with factory() as session:
        result = await _uploader(session, storage).execute(
            UploadSourceDocumentInput(
                caller=caller,
                collection_id=collection_id,
                file_name=file_name,
                content=_xlsx(rows if rows is not None else [["ZOO-1", "Jaguar"]]),
            )
        )
        await session.commit()
    return result


async def test_upload_indexes_rows_and_marks_indexed() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    result = await _upload(
        factory, storage, rows=[["ZOO-1", "Jaguar"], ["ZOO-2", "Arara"]]
    )

    assert result.created is True
    document = result.document
    assert document.status == SourceDocumentStatus.INDEXED
    assert document.row_count == 2
    assert document.indexed_at == _NOW
    async with factory() as session:
        rows = await _object_rows(session)
    assert [(r.sheet, r.row_number) for r in rows] == [("Objects", 2), ("Objects", 3)]
    assert rows[0].cells == {"Inventory No": "ZOO-1", "Name": "Jaguar"}
    assert rows[0].collection_id == str(_REPTILES)


async def test_identical_reupload_dedupes() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    first = await _upload(factory, storage)
    second = await _upload(factory, storage)

    assert second.created is False
    assert second.document.id == first.document.id
    async with factory() as session:
        assert len(await _object_rows(session)) == 1


async def test_new_version_replaces_previous_source_and_rows() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    first = await _upload(factory, storage, rows=[["ZOO-1", "Jaguar"]])
    second = await _upload(
        factory, storage, rows=[["ZOO-1", "Jaguar"], ["ZOO-9", "Tuiuiú"]]
    )

    assert second.created is True
    assert second.document.id != first.document.id
    assert second.replaced_file_reference == first.document.file_reference
    async with factory() as session:
        repo = SqlAlchemySourceDocumentRepository(session)
        old = await repo.get_by_id(first.document.id)
        assert old is not None and old.is_deleted
        live = await repo.list_live_by_collection(_REPTILES)
        assert [d.id for d in live] == [second.document.id]
        rows = await _object_rows(session)
        assert {r.source_document_id for r in rows} == {str(second.document.id)}
        assert len(rows) == 2


async def test_delete_removes_index_rows_and_soft_deletes() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    uploaded = await _upload(factory, storage)

    async with factory() as session:
        reference = await DeleteSourceDocument(
            SqlAlchemyCollectionRepository(session),
            SqlAlchemySourceDocumentRepository(session),
            SqlAlchemyCollectionObjectIndex(session),
            _FakeClock(),
        ).execute(_ADMIN, uploaded.document.id)
        await session.commit()

    assert reference == uploaded.document.file_reference
    async with factory() as session:
        assert await _object_rows(session) == []
        repo = SqlAlchemySourceDocumentRepository(session)
        record = await repo.get_by_id(uploaded.document.id)
        assert record is not None and record.is_deleted
        assert await repo.list_live_by_collection(_REPTILES) == []


async def test_reindex_rebuilds_rows_from_stored_file() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    uploaded = await _upload(factory, storage, rows=[["ZOO-1", "Jaguar"]])

    async with factory() as session:
        document = await ReindexSourceDocument(
            SqlAlchemyCollectionRepository(session),
            SqlAlchemySourceDocumentRepository(session),
            storage,
            OpenpyxlCollectionObjectParser(),
            SqlAlchemyCollectionObjectIndex(session),
            _FakeClock(),
        ).execute(_ADMIN, uploaded.document.id)
        await session.commit()

    assert document.status == SourceDocumentStatus.INDEXED
    assert document.row_count == 1
    async with factory() as session:
        assert len(await _object_rows(session)) == 1


async def test_curator_cannot_manage_unassigned_collection() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    with pytest.raises(AccessDenied):
        await _upload(factory, storage, caller=_CURATOR, collection_id=_FISH)
    assert storage.saved == {}  # nothing persisted before the scope check


async def test_assigned_curator_manages_their_collection() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    async with factory() as session:
        await AssignCollectionCurator(
            SqlAlchemyCollectionRepository(session), _FakeClock()
        ).execute(_ADMIN, _REPTILES, _CURATOR.id)
        await session.commit()

    result = await _upload(factory, storage, caller=_CURATOR)
    assert result.document.status == SourceDocumentStatus.INDEXED

    async with factory() as session:
        documents = await ListCollectionSourceDocuments(
            SqlAlchemyCollectionRepository(session),
            SqlAlchemySourceDocumentRepository(session),
        ).execute(_CURATOR, _REPTILES)
    assert [d.id for d in documents] == [result.document.id]


async def test_collections_management_manages_any_collection() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    result = await _upload(factory, storage, caller=_MANAGER, collection_id=_FISH)
    assert result.document.status == SourceDocumentStatus.INDEXED


async def test_curator_cannot_assign_curators() -> None:
    factory = await _session_factory()
    async with factory() as session:
        with pytest.raises(InsufficientGroup):
            await AssignCollectionCurator(
                SqlAlchemyCollectionRepository(session), _FakeClock()
            ).execute(_CURATOR, _REPTILES, _CURATOR.id)


async def test_assign_records_assigned_by_and_remove_revokes() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        assignment = await AssignCollectionCurator(repo, _FakeClock()).execute(
            _ADMIN, _REPTILES, _CURATOR.id
        )
        await session.commit()
    assert assignment.assigned_by == _ADMIN.id

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        # Idempotent: re-assigning returns the existing assignment.
        again = await AssignCollectionCurator(repo, _FakeClock()).execute(
            _ADMIN, _REPTILES, _CURATOR.id
        )
        assert again.assigned_by == _ADMIN.id
        await RemoveCollectionCurator(repo).execute(_ADMIN, _REPTILES, _CURATOR.id)
        await session.commit()

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        assert await repo.list_curated_collection_ids(_CURATOR.id) == set()


async def test_list_manageable_collections_flags_scope() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    await _upload(factory, storage)  # one live document in REPTILES & AMPHIBIANS
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        await AssignCollectionCurator(repo, _FakeClock()).execute(
            _ADMIN, _FISH, _CURATOR.id
        )
        await session.commit()

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        views = {
            v.collection.name: v
            for v in await ListManageableCollections(repo).execute(_CURATOR)
        }
    assert views["FISH"].manageable is True
    assert views["REPTILES & AMPHIBIANS"].manageable is False
    assert views["REPTILES & AMPHIBIANS"].document_count == 1
    assert views["REPTILES & AMPHIBIANS"].area_name == "Zoology"
    assert [c.permission_id for c in views["FISH"].curators] == [str(_CURATOR.id)]


async def test_error_status_when_stored_file_is_not_a_spreadsheet() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    uploaded = await _upload(factory, storage)
    # Corrupt the stored file, then reindex: the failure must surface as ERROR.
    storage.saved[uploaded.document.file_reference] = b"corrupted"

    async with factory() as session:
        document = await ReindexSourceDocument(
            SqlAlchemyCollectionRepository(session),
            SqlAlchemySourceDocumentRepository(session),
            storage,
            OpenpyxlCollectionObjectParser(),
            SqlAlchemyCollectionObjectIndex(session),
            _FakeClock(),
        ).execute(_ADMIN, uploaded.document.id)
        await session.commit()

    assert document.status == SourceDocumentStatus.ERROR
    assert document.error_message is not None
    async with factory() as session:
        count = await session.scalar(select(func.count(CollectionObjectRecord.id)))
    assert count == 0


# ── Collection catalog admin (SYS_ADMIN only) ───────────────────────────────


async def test_create_collection_as_sys_admin() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        collection = await CreateCollection(repo, _FakeClock()).execute(
            CreateCollectionInput(
                caller=_ADMIN, name="Mineralogy", area_id=_ZOOLOGY_AREA
            )
        )
        await session.commit()
    assert collection.created_at == _NOW
    assert collection.area_id == _ZOOLOGY_AREA

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        names = {c.name for c in await repo.list_all()}
    assert "Mineralogy" in names


async def test_create_collection_trims_whitespace_from_name() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        collection = await CreateCollection(repo, _FakeClock()).execute(
            CreateCollectionInput(
                caller=_ADMIN, name="  Mineralogy  ", area_id=_ZOOLOGY_AREA
            )
        )
        await session.commit()
    assert collection.name == "Mineralogy"

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        stored = await repo.get_by_id(collection.id)
    assert stored is not None
    assert stored.name == "Mineralogy"


async def test_get_collection_returns_enriched_view() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    await _upload(factory, storage)
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        await AssignCollectionCurator(repo, _FakeClock()).execute(
            _ADMIN, _REPTILES, _CURATOR.id
        )
        await session.commit()

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        view = await GetCollection(repo).execute(_ADMIN, _REPTILES)
    assert view.collection.name == "REPTILES & AMPHIBIANS"
    assert view.area_name == "Zoology"
    assert view.document_count == 1
    assert [c.permission_id for c in view.curators] == [str(_CURATOR.id)]
    assert view.manageable is True


async def test_get_collection_raises_not_found_for_unknown_id() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(CollectionNotFound):
            await GetCollection(repo).execute(_ADMIN, CollectionId("nope"))


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_create_collection_blocked_for_non_sys_admin(caller: Actor) -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(InsufficientGroup):
            await CreateCollection(repo, _FakeClock()).execute(
                CreateCollectionInput(
                    caller=caller, name="Mineralogy", area_id=_ZOOLOGY_AREA
                )
            )


async def test_update_collection_renames() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        await UpdateCollection(repo, _FakeClock()).execute(
            UpdateCollectionInput(
                caller=_ADMIN,
                collection_id=_REPTILES,
                name="REPTILES & AMPHIBIANS Renamed",
            )
        )
        await session.commit()

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        collection = await repo.get_by_id(_REPTILES)
    assert collection is not None
    assert collection.name == "REPTILES & AMPHIBIANS Renamed"


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_update_collection_blocked_for_non_sys_admin(caller: Actor) -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(InsufficientGroup):
            await UpdateCollection(repo, _FakeClock()).execute(
                UpdateCollectionInput(
                    caller=caller, collection_id=_REPTILES, name="Whatever"
                )
            )


async def test_remove_collection_deletes_everything() -> None:
    factory = await _session_factory()
    storage = _FakeStorage()
    uploaded = await _upload(factory, storage, rows=[["ZOO-1", "Jaguar"]])

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        await AssignCollectionCurator(repo, _FakeClock()).execute(
            _ADMIN, _REPTILES, _CURATOR.id
        )
        await session.commit()

    async with factory() as session:
        file_references = await RemoveCollection(
            SqlAlchemyCollectionRepository(session),
            SqlAlchemySourceDocumentRepository(session),
            SqlAlchemyCollectionObjectIndex(session),
        ).execute(_ADMIN, _REPTILES)
        await session.commit()

    assert file_references == [uploaded.document.file_reference]

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        # The collection itself is gone.
        assert await repo.get_by_id(_REPTILES) is None
        # FISH, untouched, is still there.
        assert (await repo.get_by_id(_FISH)) is not None
        # Its curator assignment is gone.
        assert await repo.list_curated_collection_ids(_CURATOR.id) == set()
        # Its source document row (live or not) is gone.
        documents = SqlAlchemySourceDocumentRepository(session)
        assert await documents.list_all_by_collection(_REPTILES) == []
        # Its indexed rows are gone.
        assert await _object_rows(session) == []


async def test_remove_collection_deletes_already_soft_deleted_documents_too() -> None:
    """A document soft-deleted before the collection is removed (its file
    already reclaimed) must still be purged from the database."""
    factory = await _session_factory()
    storage = _FakeStorage()
    uploaded = await _upload(factory, storage)
    async with factory() as session:
        await DeleteSourceDocument(
            SqlAlchemyCollectionRepository(session),
            SqlAlchemySourceDocumentRepository(session),
            SqlAlchemyCollectionObjectIndex(session),
            _FakeClock(),
        ).execute(_ADMIN, uploaded.document.id)
        await session.commit()

    async with factory() as session:
        await RemoveCollection(
            SqlAlchemyCollectionRepository(session),
            SqlAlchemySourceDocumentRepository(session),
            SqlAlchemyCollectionObjectIndex(session),
        ).execute(_ADMIN, _REPTILES)
        await session.commit()

    async with factory() as session:
        documents = SqlAlchemySourceDocumentRepository(session)
        assert await documents.list_all_by_collection(_REPTILES) == []


async def test_remove_collection_raises_not_found_for_unknown_id() -> None:
    factory = await _session_factory()
    async with factory() as session:
        with pytest.raises(CollectionNotFound):
            await RemoveCollection(
                SqlAlchemyCollectionRepository(session),
                SqlAlchemySourceDocumentRepository(session),
                SqlAlchemyCollectionObjectIndex(session),
            ).execute(_ADMIN, CollectionId("nope"))


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_remove_collection_blocked_for_non_sys_admin(caller: Actor) -> None:
    factory = await _session_factory()
    async with factory() as session:
        with pytest.raises(InsufficientGroup):
            await RemoveCollection(
                SqlAlchemyCollectionRepository(session),
                SqlAlchemySourceDocumentRepository(session),
                SqlAlchemyCollectionObjectIndex(session),
            ).execute(caller, _REPTILES)


async def test_list_searchable_collections_returns_full_catalogue() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        names = {c.name for c in await ListSearchableCollections(repo).execute(_ADMIN)}
    assert names == {"REPTILES & AMPHIBIANS", "FISH"}


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_collections_management_and_curator_cannot_assign_or_remove(
    caller: Actor,
) -> None:
    """Narrowed rule: only SYS_ADMIN administers curators now, not
    COLLECTIONS_MANAGEMENT (previously allowed)."""
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(InsufficientGroup):
            await AssignCollectionCurator(repo, _FakeClock()).execute(
                caller, _REPTILES, _CURATOR.id
            )
        with pytest.raises(InsufficientGroup):
            await RemoveCollectionCurator(repo).execute(caller, _REPTILES, _CURATOR.id)


class _FakePermissionReader:
    def __init__(self, views: list[PermissionView]) -> None:
        self._views = views

    async def get_detail(self, permission_id: object) -> PermissionView | None:
        raise NotImplementedError

    async def list_by_group(self, group: GroupName) -> list[PermissionView]:
        return [v for v in self._views if v.group == group]


_CURATORIAL_CANDIDATE = PermissionView(
    permission_id=str(_CURATOR.id),
    user=UserView(id="u-cur", name="Carla Curadora", email="carla@museum.pt"),
    group=GroupName.CURATORIAL,
)
_MANAGER_VIEW = PermissionView(
    permission_id=str(_MANAGER.id),
    user=UserView(id="u-mgr", name="Marco Gestor", email="marco@museum.pt"),
    group=GroupName.COLLECTIONS_MANAGEMENT,
)


async def test_list_curator_candidates_returns_curatorial_group_only() -> None:
    reader = _FakePermissionReader([_CURATORIAL_CANDIDATE, _MANAGER_VIEW])
    candidates = await ListCuratorCandidates(reader).execute(_ADMIN)
    assert [c.permission_id for c in candidates] == [str(_CURATOR.id)]


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_list_curator_candidates_blocked_for_non_sys_admin(caller: Actor) -> None:
    reader = _FakePermissionReader([_CURATORIAL_CANDIDATE])
    with pytest.raises(InsufficientGroup):
        await ListCuratorCandidates(reader).execute(caller)


# ── Collection areas (SYS_ADMIN only) ───────────────────────────────────────


async def test_create_collection_area_as_sys_admin() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        area = await CreateCollectionArea(repo, _FakeClock()).execute(
            CreateCollectionAreaInput(caller=_ADMIN, name="Documentation & Media")
        )
        await session.commit()
    assert area.created_at == _NOW

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        names = {a.name for a in await repo.list_areas()}
    assert "Documentation & Media" in names


async def test_create_collection_area_trims_whitespace_from_name() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        area = await CreateCollectionArea(repo, _FakeClock()).execute(
            CreateCollectionAreaInput(caller=_ADMIN, name="  Media  ")
        )
        await session.commit()
    assert area.name == "Media"


async def test_create_collection_area_rejects_blank_name() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(ValueError):
            await CreateCollectionArea(repo, _FakeClock()).execute(
                CreateCollectionAreaInput(caller=_ADMIN, name="   ")
            )


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_create_collection_area_blocked_for_non_sys_admin(caller: Actor) -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(InsufficientGroup):
            await CreateCollectionArea(repo, _FakeClock()).execute(
                CreateCollectionAreaInput(caller=caller, name="Media")
            )


async def test_update_collection_area_renames() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        await UpdateCollectionArea(repo, _FakeClock()).execute(
            UpdateCollectionAreaInput(
                caller=_ADMIN, area_id=_ZOOLOGY_AREA, name="Natural Sciences"
            )
        )
        await session.commit()

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        area = await repo.get_area_by_id(_ZOOLOGY_AREA)
    assert area is not None
    assert area.name == "Natural Sciences"


async def test_update_collection_area_raises_not_found_for_unknown_id() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(CollectionAreaNotFound):
            await UpdateCollectionArea(repo, _FakeClock()).execute(
                UpdateCollectionAreaInput(
                    caller=_ADMIN, area_id=CollectionAreaId("nope"), name="Whatever"
                )
            )


async def test_list_collection_areas_returns_counts() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        areas = await ListCollectionAreas(repo).execute(_ADMIN)
        views = {v.area.name: v for v in areas}
    assert views["Zoology"].collection_count == 2  # REPTILES & AMPHIBIANS + FISH
    assert views["Temporary Area"].collection_count == 0


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_list_collection_areas_blocked_for_non_sys_admin(caller: Actor) -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(InsufficientGroup):
            await ListCollectionAreas(repo).execute(caller)


async def test_remove_collection_area_blocked_while_collections_assigned() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(CollectionAreaInUse):
            await RemoveCollectionArea(repo).execute(_ADMIN, _ZOOLOGY_AREA)


async def test_remove_collection_area_succeeds_once_empty() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        await RemoveCollectionArea(repo).execute(_ADMIN, _TEMPORARY_AREA)
        await session.commit()

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        assert await repo.get_area_by_id(_TEMPORARY_AREA) is None


async def test_remove_collection_area_raises_not_found_for_unknown_id() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(CollectionAreaNotFound):
            await RemoveCollectionArea(repo).execute(_ADMIN, CollectionAreaId("nope"))


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_remove_collection_area_blocked_for_non_sys_admin(caller: Actor) -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(InsufficientGroup):
            await RemoveCollectionArea(repo).execute(caller, _TEMPORARY_AREA)


async def test_move_collection_to_area() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        collection = await MoveCollectionToArea(repo, _FakeClock()).execute(
            MoveCollectionToAreaInput(
                caller=_ADMIN, collection_id=_REPTILES, area_id=_TEMPORARY_AREA
            )
        )
        await session.commit()
    assert collection.area_id == _TEMPORARY_AREA

    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        stored = await repo.get_by_id(_REPTILES)
        counts = await repo.count_collections_by_area()
    assert stored is not None
    assert stored.area_id == _TEMPORARY_AREA
    assert counts[_TEMPORARY_AREA] == 1
    assert counts[_ZOOLOGY_AREA] == 1  # only FISH left


async def test_move_collection_to_area_raises_not_found_for_unknown_area() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(CollectionAreaNotFound):
            await MoveCollectionToArea(repo, _FakeClock()).execute(
                MoveCollectionToAreaInput(
                    caller=_ADMIN,
                    collection_id=_REPTILES,
                    area_id=CollectionAreaId("nope"),
                )
            )


async def test_move_unknown_collection_to_area_raises_not_found() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(CollectionNotFound):
            await MoveCollectionToArea(repo, _FakeClock()).execute(
                MoveCollectionToAreaInput(
                    caller=_ADMIN,
                    collection_id=CollectionId("nope"),
                    area_id=_TEMPORARY_AREA,
                )
            )


@pytest.mark.parametrize("caller", [_MANAGER, _CURATOR])
async def test_move_collection_to_area_blocked_for_non_sys_admin(caller: Actor) -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(InsufficientGroup):
            await MoveCollectionToArea(repo, _FakeClock()).execute(
                MoveCollectionToAreaInput(
                    caller=caller, collection_id=_REPTILES, area_id=_TEMPORARY_AREA
                )
            )


async def test_create_collection_raises_not_found_for_unknown_area() -> None:
    factory = await _session_factory()
    async with factory() as session:
        repo = SqlAlchemyCollectionRepository(session)
        with pytest.raises(CollectionAreaNotFound):
            await CreateCollection(repo, _FakeClock()).execute(
                CreateCollectionInput(
                    caller=_ADMIN, name="Mineralogy", area_id=CollectionAreaId("nope")
                )
            )
