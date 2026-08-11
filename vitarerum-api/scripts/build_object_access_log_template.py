"""Annotate the MUHNAC RAIS form with the placeholders the DOCX renderer fills.

The museum ships ``RAIS_formColecoesAcessoInSituRegisto_MUHNAC_2026.docx`` as a
blank form: no content controls, no merge fields, and fifteen empty rows in the
object table. This script takes that untouched form and writes Jinja
placeholders into it, producing the template asset consumed by
``app.use_of_collections.infrastructure.object_access_log_docx``.

Re-run it whenever MUHNAC revises the form, so the layout stays theirs and only
the placeholders are ours::

    uv run python scripts/build_object_access_log_template.py

Run from the repository root (or pass explicit paths); the defaults resolve
against this file's location.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

_API_ROOT = Path(__file__).resolve().parent.parent
_DEFAULT_SOURCE = (
    _API_ROOT.parent
    / "docs"
    / "templates"
    / "RAIS_formColecoesAcessoInSituRegisto_MUHNAC_2026.docx"
)
_DEFAULT_TARGET = (
    _API_ROOT
    / "app"
    / "use_of_collections"
    / "infrastructure"
    / "templates"
    / "rais_object_access_log.docx"
)

# Table/row/column coordinates of the blank form. Kept together so a revised
# form only needs these numbers checked against `--inspect` output.
_HEADER_TABLE = 0
_HEADER_ROW = 2
_IDENTIFICATION_TABLE = 1
_OBJECTS_TABLE = 2
_SIGNATURE_TABLE = 3


def _set_cell_text(cell: Any, text: str) -> None:
    """Replace a cell's text, keeping the paragraph's own run formatting.

    The blank cells carry their font on the paragraph mark (``w:pPr/w:rPr``)
    rather than on a run, so a naive ``add_run`` would render the placeholder —
    and therefore the filled value — in the document default font instead of the
    form's.
    """
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    paragraph_properties = paragraph._p.find(qn("w:pPr"))
    if paragraph_properties is not None:
        mark_properties = paragraph_properties.find(qn("w:rPr"))
        if mark_properties is not None:
            run._element.insert(0, deepcopy(mark_properties))


def _inspect(document: Any) -> None:
    for table_index, table in enumerate(document.tables):
        print(f"--- table {table_index}: {len(table.rows)}x{len(table.columns)}")
        for row_index, row in enumerate(table.rows):
            cells = [cell.text.replace("\n", "\\n") for cell in row.cells]
            print(f"   r{row_index} {cells!r}")


def annotate(source: Path, target: Path) -> None:
    document = Document(str(source))

    header = document.tables[_HEADER_TABLE].rows[_HEADER_ROW].cells
    _set_cell_text(header[0], "{{ reference_number }}")
    _set_cell_text(header[2], "{{ issued_on }}")
    _set_cell_text(header[3], "{{ requester }}")

    identification = document.tables[_IDENTIFICATION_TABLE]
    _set_cell_text(identification.rows[0].cells[1], "{{ researcher_name }}")
    _set_cell_text(identification.rows[1].cells[1], "{{ researcher_email }}")
    _set_cell_text(identification.rows[2].cells[1], "{{ collection }}")
    _set_cell_text(identification.rows[3].cells[1], "{{ curator }}")

    # Row 1 opens the loop, row 2 is the repeated body, row 3 closes it; docxtpl
    # drops the two rows holding `{%tr %}` tags. The remaining blank rows go —
    # the renderer pads the object list back up to the form's fifteen lines.
    objects = document.tables[_OBJECTS_TABLE]
    _set_cell_text(objects.rows[1].cells[0], "{%tr for object in objects %}")
    body = objects.rows[2].cells
    _set_cell_text(body[0], "{{ object.inventory_number }}")
    _set_cell_text(body[1], "{{ object.designation }}")
    _set_cell_text(body[2], "{{ object.object_type }}")
    _set_cell_text(body[3], "{{ object.number_of_objects }}")
    _set_cell_text(body[4], "{{ object.access_date }}")
    _set_cell_text(body[5], "{{ object.observations }}")
    _set_cell_text(objects.rows[3].cells[0], "{%tr endfor %}")
    for row in list(objects.rows)[4:]:
        objects._tbl.remove(row._tr)

    signature = document.tables[_SIGNATURE_TABLE]
    _set_cell_text(signature.rows[1].cells[1], "{{ conclusion_date }}")

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=_DEFAULT_SOURCE)
    parser.add_argument("--target", type=Path, default=_DEFAULT_TARGET)
    parser.add_argument(
        "--inspect",
        action="store_true",
        help="Print the source form's table coordinates and exit.",
    )
    args = parser.parse_args()

    if args.inspect:
        _inspect(Document(str(args.source)))
        return

    annotate(args.source, args.target)
    print(f"Wrote {args.target}")


if __name__ == "__main__":
    main()
